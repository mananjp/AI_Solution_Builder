"""
AI Solution Builder — Export Engine rendering services (Section 11)

Byte-level builders for PDF, DOCX, XLSX, and PPTX exports derived from a
serialized solution bundle. DOCX/XLSX exports embed a stable Field Map so
re-imported, edited files can be diffed back against live artifacts.

Each builder is pure (input: dict -> output: bytes) to stay trivially
testable and API-agnostic.
"""

import io
import json
from typing import Any


def serialize_solution(solution: Any) -> dict[str, Any]:
    """Flatten an ORM Solution + artifacts into a JSON-safe export bundle."""
    return {
        "id": str(solution.id),
        "title": solution.title,
        "description": solution.description,
        "status": solution.status,
        "created_at": solution.created_at.isoformat() if solution.created_at else None,
        "ai_state": (solution.ai_state or {}).get("analysis_summary"),
        "artifacts": [
            {
                "id": str(a.id),
                "artifact_type": a.artifact_type,
                "title": a.title,
                "version": a.version,
                "content": a.content or {},
                "content_text": a.content_text or "",
            }
            for a in solution.artifacts
        ],
    }


def field_map(bundle: dict[str, Any]) -> list[dict[str, str]]:
    """Build a flat ``field_id -> value`` index for editable exports."""
    fields: list[dict[str, str]] = []
    fields.append({"field_id": "solution.title", "value": bundle.get("title", "")})
    fields.append(
        {"field_id": "solution.description", "value": bundle.get("description", "") or ""}
    )
    for art in bundle.get("artifacts", []):
        base = f"{art['artifact_type']}.v{art['version']}."
        fields.append({"field_id": base + "title", "value": art.get("title", "")})
        for key, value in (art.get("content") or {}).items():
            fields.append(
                {
                    "field_id": base + str(key),
                    "value": json.dumps(value) if not isinstance(value, str) else value,
                }
            )
    return fields


# ── PDF ────────────────────────────────────────────
def build_pdf_bundle(bundle: dict[str, Any]) -> bytes:
    """Render a PDF architecture report from the solution bundle."""
    from xml.sax.saxutils import escape

    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    styles = getSampleStyleSheet()
    body = ParagraphStyle(name="BodyPlain", parent=styles["BodyText"], spaceAfter=6)

    doc_buf = io.BytesIO()
    doc = SimpleDocTemplate(
        doc_buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=bundle["title"],
    )

    story: list[Any] = [Paragraph(escape(bundle["title"]), styles["Title"])]
    story.append(
        Paragraph(
            f"Status: {escape((bundle.get('status') or '').upper())} — "
            f"Version bundle exported by AI Solution Builder",
            styles["Italic"],
        )
    )
    story.append(Spacer(1, 8))
    story.append(Paragraph(escape(bundle.get("description") or ""), body))

    for art in bundle.get("artifacts", []):
        story.append(PageBreak())
        story.append(Paragraph(f"{escape(art['title'])} (v{art['version']})", styles["Heading1"]))
        story.append(Paragraph(f"Type: {escape(art['artifact_type'])}", styles["Italic"]))
        story.append(Spacer(1, 6))
        if art.get("content_text"):
            for line in art["content_text"].splitlines():
                story.append(Paragraph(escape(line), body))
        else:
            rows = [
                [
                    Paragraph("<b>Field</b>", styles["BodyText"]),
                    Paragraph("<b>Value</b>", styles["BodyText"]),
                ]
            ]
            for key, value in (art.get("content") or {}).items():
                rendered = value if isinstance(value, str) else json.dumps(value)
                rows.append(
                    [
                        Paragraph(escape(str(key)), styles["BodyText"]),
                        Paragraph(escape(rendered), body),
                    ]
                )
            if len(rows) > 1:
                table = Table(rows, colWidths=[45 * mm, 105 * mm])
                table.setStyle(
                    TableStyle(
                        [
                            ("GRID", (0, 0), (-1, -1), 0.4, "#cbd5e1"),
                            ("BACKGROUND", (0, 0), (-1, 0), "#e2e8f0"),
                            ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ]
                    )
                )
                story.append(table)

    doc.build(story)
    return doc_buf.getvalue()


# ── DOCX ───────────────────────────────────────────
def build_docx_bundle(bundle: dict[str, Any]) -> bytes:
    """Render an editable DOCX with an appended Field Map (field_id -> value)."""
    from docx import Document

    document = Document()
    document.add_heading(bundle.get("title", "Solution Export"), level=0)
    document.add_paragraph(f"Status: {(bundle.get('status') or '').upper()}")
    if bundle.get("description"):
        document.add_paragraph(bundle["description"])

    for art in bundle.get("artifacts", []):
        document.add_heading(f"{art['title']} (v{art['version']})", level=1)
        document.add_paragraph(f"Type: {art['artifact_type']}")
        if art.get("content_text"):
            for line in art["content_text"].splitlines():
                document.add_paragraph(line)
        else:
            for key, value in (art.get("content") or {}).items():
                rendered = value if isinstance(value, str) else json.dumps(value)
                document.add_paragraph(f"{key}: {rendered}")

    # Editable Field Map — consumers may re-import and diff on field_id.
    document.add_heading("Field Map (editable)", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    header = table.rows[0].cells
    header[0].text = "field_id"
    header[1].text = "value"
    for entry in field_map(bundle):
        cells = table.add_row().cells
        cells[0].text = entry["field_id"]
        cells[1].text = entry["value"]

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ── XLSX ───────────────────────────────────────────
def build_xlsx_bundle(bundle: dict[str, Any]) -> bytes:
    """Render one worksheet per artifact plus a Field Map sheet."""
    from openpyxl import Workbook

    wb = Workbook()

    for art in bundle.get("artifacts", []):
        ws = wb.create_sheet(title=(art["artifact_type"].replace("_", "-"))[:31])
        ws.append(["field", "value"])
        if art.get("content_text"):
            for idx, line in enumerate(art["content_text"].splitlines()):
                ws.append([f"line-{idx + 1}", line])
        else:
            for key, value in (art.get("content") or {}).items():
                ws.append([key, json.dumps(value) if not isinstance(value, str) else value])

    field_sheet = wb.create_sheet(title="Field-Map")
    field_sheet.append(["field_id", "value"])
    for entry in field_map(bundle):
        field_sheet.append([entry["field_id"], entry["value"]])

    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── PPTX ───────────────────────────────────────────
def build_pptx_bundle(bundle: dict[str, Any]) -> bytes:
    """Render a slide deck: title, summary, and one slide per artifact."""
    from pptx import Presentation
    from pptx.util import Pt

    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = bundle.get("title", "Solution Export")
    if bundle.get("description"):
        title_slide.placeholders[1].text = bundle["description"]

    for art in bundle.get("artifacts", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"{art['title']} (v{art['version']})"
        body = slide.placeholders[1].text_frame
        lines = (
            art.get("content_text", "").splitlines()
            if art.get("content_text")
            else [
                f"{key}: {value if isinstance(value, str) else json.dumps(value)}"
                for key, value in (art.get("content") or {}).items()
            ]
        )
        for idx, line in enumerate(lines[:8]):
            paragraph = body.paragraphs[0] if idx == 0 else body.add_paragraph()
            paragraph.text = line
            paragraph.font.size = Pt(14)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
