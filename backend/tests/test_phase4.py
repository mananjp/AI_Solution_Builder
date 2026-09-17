"""Phase 4 tests: multilingual pipeline, rich exports, Figma/deployer, collaboration, RLS."""

import io
import json
from uuid import uuid4

import pytest

from app.core.i18n import (
    detect_language,
    normalize_language_code,
    pick_best_language,
    translate_text,
)
from app.services.deployer import DeployError, build_repo_files, deploy_to_github
from app.services.exporters import (
    build_docx_bundle,
    build_pdf_bundle,
    build_pptx_bundle,
    build_xlsx_bundle,
    field_map,
    serialize_solution,
)
from app.services.figma import build_figma_manifest
from app.services.rls import (
    build_policy_sql,
    enable_org_rls,
    set_request_org,
    tables_with_org_column,
)


# ── Multilingual / i18n ─────────────────────────────────────────────
class TestI18n:
    def test_detect_language_english_and_spanish(self):
        assert detect_language("This is a simple English sentence for the analysis.", "en") == "en"
        assert (
            detect_language("Esto es una frase sencilla en español para el análisis.", "en") == "es"
        )

    def test_detect_language_empty_falls_back(self):
        assert detect_language("", "en") == "en"
        assert detect_language("   \n\t ", "de") == "de"

    def test_normalize_language_codes(self):
        assert normalize_language_code("fr-FR") == "fr"
        assert normalize_language_code("zh-CN") == "zh"
        assert normalize_language_code("EN") == "en"
        assert normalize_language_code("xx") == "en"  # unsupported -> default
        assert normalize_language_code("") == "en"

    def test_pick_best_language_precedence(self):
        # Explicit header wins over Accept-Language
        assert pick_best_language("es", "de", "English source") == "de"
        # Accept-Language hint used when no explicit header
        assert pick_best_language("fr-FR", "", "English source") == "fr"
        # Falls back to detection otherwise
        assert pick_best_language("", "", "Esto es un texto en español") == "es"

    async def test_translate_text_mock_is_passthrough(self):
        # Under the mock provider translation is a no-op (deterministic tests)
        assert await translate_text("Some English text", "es") == "Some English text"
        assert await translate_text("", "es") == ""
        assert await translate_text("Text", "en") == "Text"


# ── Rich exports ────────────────────────────────────────────────────
def _sample_bundle() -> dict:
    return {
        "id": str(uuid4()),
        "title": "CRM Platform",
        "description": "A lead management platform",
        "status": "complete",
        "created_at": "2026-01-01T00:00:00+00:00",
        "ai_state": "summary",
        "artifacts": [
            {
                "id": str(uuid4()),
                "artifact_type": "hld",
                "title": "High-Level Design",
                "version": 1,
                "content": {"system_overview": "modular", "components": ["web", "api"]},
                "content_text": "",
            },
            {
                "id": str(uuid4()),
                "artifact_type": "wireframe",
                "title": "Leads Dashboard",
                "version": 1,
                "content": {
                    "module": "crm",
                    "screens": [
                        {
                            "name": "Leads",
                            "components": [
                                {
                                    "title": "Table",
                                    "type": "data_table",
                                    "description": "Rows of leads",
                                }
                            ],
                        }
                    ],
                },
                "content_text": "",
            },
        ],
    }


class TestExportBuilders:
    def test_serialize_solution(self):
        class FakeArtifact:
            id = uuid4()
            artifact_type = "hld"
            title = "HLD"
            version = 3
            content = {"k": "v"}
            content_text = "text"

        class FakeSolution:
            id = uuid4()
            title = "Sol"
            description = "desc"
            status = "complete"
            created_at = None
            artifacts = [FakeArtifact()]
            ai_state = {"analysis_summary": "sum"}

        bundle = serialize_solution(FakeSolution())
        assert bundle["title"] == "Sol"
        assert bundle["artifacts"][0]["version"] == 3

    def test_field_map_contains_stable_ids(self):
        fields = field_map(_sample_bundle())
        ids = {f["field_id"] for f in fields}
        assert "solution.title" in ids
        assert any("hld.v1." in i for i in ids)

    def test_pdf_bytes(self):
        data = build_pdf_bundle(_sample_bundle())
        assert data.startswith(b"%PDF")

    def test_docx_bytes_and_field_map(self):
        from docx import Document

        data = build_docx_bundle(_sample_bundle())
        assert data.startswith(b"PK")
        doc = Document(io.BytesIO(data))
        # Heading-level title present and field map table exists
        assert len(doc.tables) == 1
        assert doc.tables[0].rows[0].cells[0].text == "field_id"

    def test_xlsx_bytes(self):
        from openpyxl import load_workbook

        data = build_xlsx_bundle(_sample_bundle())
        assert data.startswith(b"PK")
        wb = load_workbook(io.BytesIO(data))
        assert "Field-Map" in wb.sheetnames

    def test_pptx_bytes(self):
        from pptx import Presentation

        data = build_pptx_bundle(_sample_bundle())
        assert data.startswith(b"PK")
        prs = Presentation(io.BytesIO(data))
        assert len(prs.slides) == 3  # title + 2 artifact slides


# ── Figma exporter ──────────────────────────────────────────────────
class TestFigma:
    def test_build_figma_manifest(self):
        manifest = build_figma_manifest(_sample_bundle())
        assert manifest["type"] == "FILE"
        assert manifest["children"][0]["type"] == "CANVAS"
        assert manifest["children"][0]["children"][0]["name"] == "Leads"

    def test_empty_wireframes(self):
        bundle = _sample_bundle()
        bundle["artifacts"] = [a for a in bundle["artifacts"] if a["artifact_type"] != "wireframe"]
        manifest = build_figma_manifest(bundle)
        assert manifest["children"]  # fallback empty canvas


# ── One-Click Deployer ──────────────────────────────────────────────
class TestDeployer:
    def test_build_repo_files(self):
        files = build_repo_files(_sample_bundle(), "name: ci")
        assert set(files) == {
            "README.md",
            "docker-compose.yml",
            "init.sql",
            "api_spec.json",
            ".github/workflows/deploy.yml",
        }

    async def test_deploy_requires_token(self):
        with pytest.raises(DeployError, match="GITHUB_TOKEN"):
            await deploy_to_github("", "some-repo", {"README.md": "# hi"})

    async def test_deploy_creates_repo(self, monkeypatch):
        monkeypatch.setattr(
            "app.services.deployer.httpx.AsyncClient", lambda *a, **k: _FakeGithubClient()
        )
        result = await deploy_to_github(
            "ghp_fake", "demo-app", build_repo_files(_sample_bundle(), "name: ci")
        )
        assert result["url"] == "https://github.com/testowner/demo-app"
        assert result["branch"] == "main"

    async def test_deploy_repo_create_failure(self, monkeypatch):
        monkeypatch.setattr(
            "app.services.deployer.httpx.AsyncClient",
            lambda *a, **k: _FakeGithubClient(fail_create=True),
        )
        with pytest.raises(DeployError, match="Repo create failed"):
            await deploy_to_github(
                "ghp_fake", "demo-app", build_repo_files(_sample_bundle(), "name: ci")
            )


class _FakeGithubClient:
    """Simplified httpx client double returning scripted GitHub responses."""

    def __init__(self, fail_create: bool = False) -> None:
        self.calls: list[tuple[str, dict]] = []
        self.fail_create = fail_create

    async def __aenter__(self):
        return self

    async def __aexit__(self, *exc):
        return None

    async def post(self, url, headers=None, json=None, **kwargs):
        self.calls.append(("post", {"url": url, "json": json}))
        return _FakeResponse(
            500 if self.fail_create else 201,
            {"owner": {"login": "testowner"}, "default_branch": "main"},
        )

    async def put(self, url, headers=None, json=None, **kwargs):
        self.calls.append(("put", {"url": url, "json": json}))
        return _FakeResponse(201, {})

    async def get(self, url, headers=None, **kwargs):
        self.calls.append(("get", {"url": url}))
        return _FakeResponse(404, {})


class _FakeResponse:
    def __init__(self, status_code: int, payload: dict) -> None:
        self.status_code = status_code
        self._payload = payload
        self.text = json.dumps(payload)

    def json(self) -> dict:
        return self._payload


# ── Row-Level Security ──────────────────────────────────────────────
class TestRLS:
    def test_tables_with_org_column(self):
        cols = [("leads", "org_id"), ("leads", "title"), ("projects", "org_id")]
        assert tables_with_org_column(cols) == ["leads", "projects"]

    def test_build_policy_sql(self):
        statements = build_policy_sql("workable_x", "leads")
        joined = "\n".join(statements)
        assert "ENABLE ROW LEVEL SECURITY" in joined
        assert 'CREATE POLICY "org_isolation"' in joined
        assert "app.org_id" in joined

    async def test_enable_org_rls(self):
        conn = _FakeConn([("leads", "org_id"), ("leads", "title")])
        await enable_org_rls(conn, "workable_x", uuid4())
        executed = "\n".join(str(s).lower() for s, _ in conn.calls)
        assert "enable row level security" in executed
        assert "create policy" in executed

    async def test_set_request_org(self):
        conn = _FakeConn([])
        await set_request_org(conn, uuid4())
        assert any("set_config('app.org_id'" in str(s) for s, _ in conn.calls)


class _FakeConn:
    def __init__(self, columns_rows: list[tuple[str, str]]) -> None:
        self.calls: list[tuple] = []
        self._rows = columns_rows

    async def execute(self, statement, params=None):
        self.calls.append((statement, params))
        result = _FakeResult(self._rows)
        self._rows = []  # only the introspection query returns rows
        return result


class _FakeResult:
    def __init__(self, rows: list) -> None:
        self._rows = rows

    def all(self):
        return list(self._rows)


# ── Integration: multilingual chat headers ─────────────────────────
class TestLanguageMiddleware:
    async def test_accept_language_sets_header(self, workspace_solution):
        resp = await workspace_solution["client"].get(
            "/",
            headers={**workspace_solution["headers"], "Accept-Language": "es-MX"},
        )
        assert resp.status_code == 200
        assert resp.headers["x-content-language"] == "es"

    async def test_explicit_content_language_wins(self, workspace_solution):
        resp = await workspace_solution["client"].get(
            "/",
            headers={
                **workspace_solution["headers"],
                "X-Content-Language": "fr",
                "Accept-Language": "de",
            },
        )
        assert resp.headers["x-content-language"] == "fr"

    async def test_missing_headers_defaults_to_en(self, workspace_solution):
        resp = await workspace_solution["client"].get("/", headers=workspace_solution["headers"])
        assert resp.headers["x-content-language"] == "en"

    async def test_mock_provider_keeps_english_reply(self, workspace_solution):
        # Translation is a no-op under the mock LLM, but the request flows
        # through the language-aware path without error.
        client = workspace_solution["client"]
        headers = {**workspace_solution["headers"], "X-Content-Language": "es"}
        resp = await client.post(
            "/api/v1/chat/send",
            json={
                "solution_id": workspace_solution["solution_id"],
                "message": "We run a B2B SaaS CRM and need lead tracking",
            },
            headers=headers,
        )
        assert resp.status_code == 200, resp.text
        assert resp.headers["x-content-language"] == "es"


# ── Integration: rich exports & Figma ──────────────────────────────
class TestRichExportEndpoints:
    @pytest.fixture(autouse=True)
    async def _generated(self, workspace_solution):
        self.ctx = workspace_solution
        resp = await workspace_solution["client"].post(
            "/api/v1/chat/send",
            json={
                "solution_id": workspace_solution["solution_id"],
                "message": "We run a B2B SaaS company selling CRM software and need lead tracking",
            },
            headers=workspace_solution["headers"],
        )
        assert resp.status_code == 200, resp.text

    async def test_pdf(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/export/{self.ctx['solution_id']}/pdf", headers=self.ctx["headers"]
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert resp.content.startswith(b"%PDF")

    async def test_docx(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/export/{self.ctx['solution_id']}/docx", headers=self.ctx["headers"]
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        assert resp.content.startswith(b"PK")

    async def test_xlsx(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/export/{self.ctx['solution_id']}/xlsx", headers=self.ctx["headers"]
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        assert resp.content.startswith(b"PK")

    async def test_pptx(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/export/{self.ctx['solution_id']}/pptx", headers=self.ctx["headers"]
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument"
        )
        assert resp.content.startswith(b"PK")

    async def test_figma_manifest(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/export/{self.ctx['solution_id']}/figma", headers=self.ctx["headers"]
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/json")
        assert json.loads(resp.text)["type"] == "FILE"

    async def test_unknown_solution_404(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/export/{uuid4()}/pdf", headers=self.ctx["headers"]
        )
        assert resp.status_code == 404


# ── Integration: one-click deploy ──────────────────────────────────
class TestDeployEndpoints:
    @pytest.fixture(autouse=True)
    async def _generated(self, workspace_solution):
        self.ctx = workspace_solution
        resp = await workspace_solution["client"].post(
            "/api/v1/chat/send",
            json={
                "solution_id": workspace_solution["solution_id"],
                "message": "We run a B2B SaaS CRM and need a deployable system",
            },
            headers=workspace_solution["headers"],
        )
        assert resp.status_code == 200, resp.text

    async def test_deploy_success(self, monkeypatch):
        async def fake_deploy(*args, **kwargs):
            return {
                "owner": "testowner",
                "repo": "app-repo",
                "url": "https://github.com/testowner/app-repo",
                "default_branch": "main",
                "branch": "main",
            }

        monkeypatch.setattr("app.api.export.deploy_to_github", fake_deploy)
        resp = await self.ctx["client"].post(
            f"/api/v1/export/{self.ctx['solution_id']}/deploy",
            json={"repo_name": "app-repo", "description": "Generated solution", "private": True},
            headers=self.ctx["headers"],
        )
        assert resp.status_code == 200, resp.text
        assert resp.json()["status"] == "deployed"
        assert resp.json()["url"].startswith("https://github.com/")

    async def test_deploy_invalid_repo_name(self):
        resp = await self.ctx["client"].post(
            f"/api/v1/export/{self.ctx['solution_id']}/deploy",
            json={"repo_name": "../escape"},
            headers=self.ctx["headers"],
        )
        assert resp.status_code == 400

    async def test_deploy_upstream_failure(self, monkeypatch):
        async def failing_deploy(*args, **kwargs):
            raise DeployError("Repo create failed")

        monkeypatch.setattr("app.api.export.deploy_to_github", failing_deploy)
        resp = await self.ctx["client"].post(
            f"/api/v1/export/{self.ctx['solution_id']}/deploy",
            json={"repo_name": "app-repo"},
            headers=self.ctx["headers"],
        )
        assert resp.status_code == 502

    async def test_deploy_reads_configured_token(self, monkeypatch):
        seen = {}

        async def capture_deploy(token, repo_name, files, description="", private=False):
            seen["token"] = token
            return {"owner": "o", "repo": repo_name, "url": "u", "branch": "main"}

        from app.core.config import settings

        monkeypatch.setattr(settings, "GITHUB_TOKEN", "ghp_fake_token")
        monkeypatch.setattr("app.api.export.deploy_to_github", capture_deploy)
        resp = await self.ctx["client"].post(
            f"/api/v1/export/{self.ctx['solution_id']}/deploy",
            json={"repo_name": "app-repo"},
            headers=self.ctx["headers"],
        )
        assert resp.status_code == 200
        assert seen["token"] == "ghp_fake_token"


# ── Integration: collaboration comments & activity ─────────────────
class TestCollaborationEndpoints:
    async def _artifact_id(self):
        resp = await self.ctx["client"].get(
            f"/api/v1/solutions/{self.ctx['solution_id']}", headers=self.ctx["headers"]
        )
        artifacts = resp.json()["artifacts"]
        return artifacts[0]["id"]

    @pytest.fixture(autouse=True)
    async def _generated(self, workspace_solution):
        self.ctx = workspace_solution
        resp = await workspace_solution["client"].post(
            "/api/v1/chat/send",
            json={
                "solution_id": workspace_solution["solution_id"],
                "message": "We run a B2B SaaS CRM and need a system",
            },
            headers=workspace_solution["headers"],
        )
        assert resp.status_code == 200, resp.text

    async def test_comment_roundtrip(self):
        artifact_id = await self._artifact_id()
        client, headers, sol_id = self.ctx["client"], self.ctx["headers"], self.ctx["solution_id"]

        created = await client.post(
            f"/api/v1/artifacts/{artifact_id}/comments",
            json={"body": "Please add SSO to this design"},
            headers=headers,
        )
        assert created.status_code == 201, created.text
        assert created.json()["body"] == "Please add SSO to this design"
        assert created.json()["resolved"] is False
        comment_id = created.json()["id"]

        listed = await client.get(f"/api/v1/artifacts/{artifact_id}/comments", headers=headers)
        assert listed.status_code == 200
        assert any(c["id"] == comment_id for c in listed.json())

        resolved = await client.patch(
            f"/api/v1/artifacts/comments/{comment_id}/resolve", headers=headers
        )
        assert resolved.status_code == 200
        assert resolved.json()["resolved"] is True

        activity = await client.get(f"/api/v1/artifacts/activity/{sol_id}", headers=headers)
        assert activity.status_code == 200
        events = activity.json()["events"]
        assert any(e["id"] == comment_id and e["artifact_type"] for e in events)

    async def test_empty_comment_rejected(self):
        artifact_id = await self._artifact_id()
        resp = await self.ctx["client"].post(
            f"/api/v1/artifacts/{artifact_id}/comments",
            json={"body": "   "},
            headers=self.ctx["headers"],
        )
        assert resp.status_code == 400

    async def test_cross_org_is_isolated(self):
        artifact_id = await self._artifact_id()
        email = f"test-{uuid4().hex[:10]}@example.com"
        resp = await self.ctx["client"].post(
            "/api/v1/auth/register",
            json={
                "email": email,
                "full_name": "Other User",
                "password": "Testpass123!",
                "org_name": f"Org-{uuid4().hex[:6]}",
            },
        )
        assert resp.status_code == 201, resp.text
        other_headers = {"Authorization": f"Bearer {resp.json()['access_token']}"}
        resp = await self.ctx["client"].post(
            f"/api/v1/artifacts/{artifact_id}/comments",
            json={"body": "trespass"},
            headers=other_headers,
        )
        assert resp.status_code == 404
        resp = await self.ctx["client"].get(
            f"/api/v1/artifacts/{artifact_id}/comments", headers=other_headers
        )
        assert resp.status_code == 404

    async def test_unknown_comment_404(self):
        resp = await self.ctx["client"].patch(
            f"/api/v1/artifacts/comments/{uuid4()}/resolve", headers=self.ctx["headers"]
        )
        assert resp.status_code == 404
